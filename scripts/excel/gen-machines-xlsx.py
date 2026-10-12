#!/usr/bin/env python3
"""Génère docs/Machines-abaques.xlsx : disponibilité des abaques par machine.

Sources : data/charts/<slug>/machine.json (dossiers ne commençant pas par « _ »)
et data/machines-sans-abaque.json. Usage : npm run machines:xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "Machines-abaques.xlsx"

TYPES = {
    "grue_mobile": "Grue mobile sur pneus",
    "grue_mobile_chenilles": "Grue mobile sur chenilles",
    "grue_tour": "Grue à tour",
    "grue_auxiliaire": "Grue auxiliaire (camion-bras)",
    "telescopique_fixe": "Chariot télescopique fixe",
    "telescopique_rotatif": "Chariot télescopique rotatif",
    "pelle_pneus": "Pelle sur pneus",
    "pelle_chenilles": "Pelle sur chenilles",
    "pelle_rail_route": "Pelle rail-route",
}
BLEU, GRIS, VERT, ROUGE, JAUNE, BORD = "004489", "F0F0F0", "7AA536", "E20025", "F2AB1B", "DCDCDC"
ARIAL = "Arial"
MENTION = ("Outil d'aide à la rédaction : l'abaque doit être vérifié sur la documentation "
           "de la machine réellement présente sur site, et validé par une personne compétente.")

# ---------------------------------------------------------------- données
machines = []
for p in sorted((ROOT / "data" / "charts").glob("*/machine.json")):
    if p.parent.name.startswith("_"):
        continue
    d = json.loads(p.read_text(encoding="utf-8"))
    src = d.get("source") or {}
    titre = (src.get("document_title") or "").strip()
    editeur = (src.get("publisher") or "").strip()
    source = " — ".join(x for x in (titre, editeur) if x) or "—"
    statut = {"verifie": "Vérifié", "a_verifier": "À vérifier"}.get(d.get("status"), "—")
    machines.append(dict(constructeur=d["manufacturer"], modele=d["model"], type=d["type_code"],
                         abaque="Oui", nb=len(d.get("configurations", [])), statut=statut,
                         source=source, doc="—"))
sans = json.loads((ROOT / "data" / "machines-sans-abaque.json").read_text(encoding="utf-8"))["machines"]
for m in sans:
    machines.append(dict(constructeur=m["manufacturer"], modele=m["model"], type=m["type_code"],
                         abaque="Non", nb=0, statut="—", source=m.get("motif") or "—",
                         doc=m.get("besoin") or "—"))
for m in machines:
    m["label"] = TYPES.get(m["type"], m["type"])
    m["nom"] = f'{m["constructeur"]} {m["modele"]}'
order = list(TYPES)
machines.sort(key=lambda m: (order.index(m["type"]) if m["type"] in order else 99,
                             m["constructeur"].lower(), m["modele"].lower()))
noms = [m["nom"] for m in machines]
assert len(set(noms)) == len(noms), "noms de machines en double"
N = len(machines) + 1  # dernière ligne de la feuille Machines

# ---------------------------------------------------------------- styles
fb = lambda **k: Font(name=ARIAL, size=k.pop("size", 10), **k)
side = Side(style="thin", color=BORD)
border = Border(left=side, right=side, top=side, bottom=side)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)


def entete(ws, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.border = fb(bold=True, color="FFFFFF"), fill(BLEU), border
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def corps(c, r, align=wrap):
    c.font, c.border, c.alignment = fb(), border, align
    c.fill = fill("FFFFFF" if r % 2 else GRIS)


def regles_ouinon(ws, rng, a_verifier=False):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Oui"'], fill=fill(VERT), font=fb(bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Non"'], fill=fill(ROUGE), font=fb(bold=True, color="FFFFFF")))
    if a_verifier:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"À vérifier"'], fill=fill(JAUNE), font=fb(bold=True, color="000000")))


wb = Workbook()
wr = wb.active
wr.title = "Recherche"
wm = wb.create_sheet("Machines")
ws = wb.create_sheet("Synthèse")

# ---------------------------------------------------------------- Machines
HEAD = ["Machine", "Constructeur", "Modèle", "Type", "Abaque disponible", "Nb configurations",
        "Statut", "Source / motif", "Document nécessaire"]
entete(wm, 1, HEAD)
for r, m in enumerate(machines, 2):
    vals = [m["nom"], m["constructeur"], m["modele"], m["label"], m["abaque"], m["nb"],
            m["statut"], m["source"], m["doc"]]
    for ci, v in enumerate(vals, 1):
        c = wm.cell(r, ci, v)
        corps(c, r, center if ci in (5, 6, 7) else wrap)
for col, w in zip("ABCDEFGHI", (30, 16, 24, 28, 14, 14, 13, 70, 45)):
    wm.column_dimensions[col].width = w
wm.row_dimensions[1].height = 32
wm.freeze_panes = "A2"
wm.auto_filter.ref = f"A1:I{N}"
regles_ouinon(wm, f"E2:E{N}")
wm.conditional_formatting.add(f"G2:G{N}", CellIsRule(operator="equal", formula=['"À vérifier"'], fill=fill(JAUNE), font=fb(bold=True)))
wb.defined_names["ListeMachines"] = DefinedName("ListeMachines", attr_text=f"Machines!$A$2:$A${N}")

# ---------------------------------------------------------------- Recherche
wr.sheet_view.showGridLines = False
wr.column_dimensions["A"].width = 30
wr.column_dimensions["B"].width = 80
wr["A1"] = "Adéquation de levage — disponibilité des abaques"
wr["A1"].font = fb(bold=True, size=16, color=BLEU)
wr["A2"] = "Choisissez une machine dans la liste déroulante (cellule bleutée ci-dessous) : les informations s'affichent automatiquement."
wr["A2"].font = fb(italic=True)
wr["A4"] = "Machine"
wr["A4"].font, wr["A4"].fill, wr["A4"].border = fb(bold=True, color="FFFFFF"), fill(BLEU), border
wr["B4"] = noms[0]
wr["B4"].font, wr["B4"].border = fb(bold=True, size=12, color=BLEU), Border(*(Side(style="medium", color=BLEU),) * 4)
wr["B4"].fill = fill(GRIS)
wr.row_dimensions[4].height = 24
dv = DataValidation(type="list", formula1="=ListeMachines", allow_blank=False, showDropDown=False,
                    showErrorMessage=True, errorTitle="Machine inconnue",
                    error="Choisissez une machine dans la liste.")
wr.add_data_validation(dv)
dv.add("B4")

def look(col):
    return f'IFERROR(INDEX(Machines!${col}$2:${col}${N},MATCH($B$4,Machines!$A$2:$A${N},0)),"—")'

lignes = [("Type", "D"), ("Abaque disponible", "E"), ("Nb de configurations", "F"),
          ("Statut", "G"), ("Source ou motif", "H"), ("Document nécessaire", "I")]
for i, (lab, col) in enumerate(lignes):
    r = 6 + i
    a = wr.cell(r, 1, lab)
    a.font, a.fill, a.border, a.alignment = fb(bold=True), fill(GRIS), border, wrap
    b = wr.cell(r, 2, "=" + look(col))
    b.font, b.border = fb(), border
    b.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    b.fill = fill("FFFFFF")
wr.cell(7, 2).font = fb(bold=True)
regles_ouinon(wr, "B7")
wr.conditional_formatting.add("B9", CellIsRule(operator="equal", formula=['"À vérifier"'], fill=fill(JAUNE), font=fb(bold=True)))
for r, h in ((10, 60), (11, 45)):
    wr.row_dimensions[r].height = h
wr["A13"] = MENTION
wr["A13"].font = fb(bold=True, color=ROUGE)
wr["A13"].alignment = wrap
wr.merge_cells("A13:B13")
wr.row_dimensions[13].height = 32
wr["A14"] = ("Légende : Oui = abaque présent dans le catalogue ; Non = abaque manquant (voir « Document nécessaire »). "
             "Statut « À vérifier » = abaque saisi mais non encore relu contre la documentation constructeur.")
wr["A14"].font = fb(italic=True)
wr["A14"].alignment = wrap
wr.merge_cells("A14:B14")
wr.row_dimensions[14].height = 32

# ---------------------------------------------------------------- Synthèse
for col, w in zip("ABCD", (32, 14, 16, 16)):
    ws.column_dimensions[col].width = w
ws.column_dimensions["E"].width = 4
ws["A1"] = "Synthèse des abaques par type de machine"
ws["A1"].font = fb(bold=True, size=14, color=BLEU)
entete(ws, 3, ["Type", "Total", "Avec abaque", "Sans abaque"])
utilises = [t for t in TYPES if any(m["type"] == t for m in machines)]
for i, t in enumerate(utilises):
    r = 4 + i
    ws.cell(r, 1, TYPES[t])
    ws.cell(r, 2, f'=COUNTIFS(Machines!$D$2:$D${N},$A{r})')
    ws.cell(r, 3, f'=COUNTIFS(Machines!$D$2:$D${N},$A{r},Machines!$E$2:$E${N},"Oui")')
    ws.cell(r, 4, f'=COUNTIFS(Machines!$D$2:$D${N},$A{r},Machines!$E$2:$E${N},"Non")')
    for ci in range(1, 5):
        corps(ws.cell(r, ci), r, center if ci > 1 else wrap)
tr = 4 + len(utilises)
ws.cell(tr, 1, "Total")
for ci, L in ((2, "B"), (3, "C"), (4, "D")):
    ws.cell(tr, ci, f"=SUM({L}4:{L}{tr - 1})")
for ci in range(1, 5):
    c = ws.cell(tr, ci)
    c.font, c.fill, c.border = fb(bold=True, color="FFFFFF"), fill(BLEU), border
    c.alignment = center if ci > 1 else wrap

r0 = tr + 3
ws.cell(r0 - 1, 1, "Machines sans abaque").font = fb(bold=True, size=12, color=BLEU)
entete(ws, r0, ["Machine", "Type", "", ""])
ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=4)
manquantes = [m for m in machines if m["abaque"] == "Non"]
for i, m in enumerate(manquantes):
    r = r0 + 1 + i
    ws.cell(r, 1, m["nom"])
    ws.cell(r, 2, m["label"])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for ci in range(1, 5):
        corps(ws.cell(r, ci), r)
ws.cell(r0 + len(manquantes) + 2, 1, "Liste des manquants figée à la génération (npm run machines:xlsx) ; les comptages ci-dessus sont des formules.").font = fb(italic=True)

for sh in (wr, wm, ws):
    sh.sheet_properties.tabColor = BLEU
wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"{OUT} : {sum(m['abaque'] == 'Oui' for m in machines)} avec abaque, {len(manquantes)} sans, {len(machines)} au total")
